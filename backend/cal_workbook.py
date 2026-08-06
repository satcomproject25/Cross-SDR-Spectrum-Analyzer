"""Calibration workbook: template generation, parsing, and solving.

One workbook, four sheets:

    README     instructions, no inputs
    SETUP      device identity and fixed RF path losses
    FREQ_CAL   known frequency vs displayed frequency
    POWER_CAL  known input power vs displayed dBFS

Design note that matters more than any other in this file
---------------------------------------------------------
Every derived quantity is recomputed here from the RAW INPUT CELLS. The
formulas written into the template are operator feedback only and their cached
results are never read. A workbook edited in Excel, LibreOffice, Google Sheets,
or by hand therefore parses identically, and a stale cached value cannot
silently corrupt a calibration.

Measurement model
-----------------
    P_in_dBm  = SigGen_dBm - attenuator_dB - cable_loss_dB
    Offset_dB = P_in_dBm - Measured_dBFS
    Base_dB   = Offset_dB + Gain_dB          (gain-normalised, device only)

Inverse, applied at runtime by backend.calibration.PowerCalibration:

    dBm = dBFS + Base(f) + delta(g) - g + external_attenuation_db
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Acceptance limits -----------------------------------------------------
CLIP_DBFS = -3.0        # above this the ADC is compressing; row is invalid
WEAK_DBFS = -100.0      # below this the tone is in the noise; row is invalid
FREQ_RESIDUAL_FLOOR_HZ = 200.0   # non-linearity threshold, absolute
FREQ_RESIDUAL_FRACTION = 0.20    # non-linearity threshold, relative to span
POWER_SCATTER_WARN_DB = 1.5      # repeat-point disagreement worth flagging

FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")   # yellow: you fill in
FILL_FIXED = PatternFill("solid", fgColor="D9D9D9")   # grey: do not change
FILL_CALC = PatternFill("solid", fgColor="E2EFDA")    # green: computed
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FONT_HEAD = Font(color="FFFFFF", bold=True, size=10)
FONT_TITLE = Font(bold=True, size=13, color="1F3864")
THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DEFAULT_FREQ_POINTS_MHZ = [70, 200, 500, 900, 1400, 2000, 2450, 3000, 4000, 5000]
DEFAULT_POWER_POINTS_MHZ = [
    70, 100, 200, 400, 600, 800, 1000, 1400, 1800, 2000,
    2200, 2450, 2700, 3000, 3500, 4000, 4500, 5000, 5500, 5800,
]


# ===========================================================================
# Template generation
# ===========================================================================
def _write_header(ws, row: int, labels: list[str]) -> None:
    for col, text in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = FILL_HEAD
        cell.font = FONT_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _widths(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _readme(wb: Workbook) -> None:
    ws = wb.create_sheet("README", 0)
    _widths(ws, {1: 4, 2: 108})
    lines = [
        (FONT_TITLE, "SDR CALIBRATION WORKBOOK"),
        (None, ""),
        (None, "Fill the YELLOW cells only. Grey cells are fixed; green cells compute themselves."),
        (None, "Upload this same file back into the application when you are done."),
        (None, ""),
        (FONT_TITLE, "BEFORE YOU START"),
        (None, "1. Warm the SDR for 15-20 minutes. Cold boards drift by 1-2 dB and several kHz."),
        (None, "2. Fit a 6-10 dB attenuator pad directly at the SDR SMA input."),
        (None, "   A pad absorbs cable reflections twice, so a 10 dB pad suppresses standing-wave"),
        (None, "   ripple by about 20 dB. Without it you freeze YOUR cable's ripple into the numbers"),
        (None, "   and the calibration is wrong by +/-2 dB the moment you connect anything else."),
        (None, "3. Enter the pad and cable loss on the SETUP sheet. Everything else is derived."),
        (None, "4. Keep LNA, AMP, sample rate, span and cabling IDENTICAL for every row."),
        (None, "5. Place the test tone 100-300 kHz off centre, never at centre: the DC blocker and"),
        (None, "   the direct-conversion LO artefact both live in the centre bin."),
        (None, ""),
        (FONT_TITLE, "SHEET BY SHEET"),
        (None, "SETUP      Device type, serial, gain stages, pad and cable loss. Required."),
        (None, "FREQ_CAL   Tune to the frequency in column A, read the MARKER, type it in column B."),
        (None, "           At least 4 rows, spread as widely as the hardware allows. 6+ is better."),
        (None, "           Widely spaced points are what separate a fixed offset from a ppm error;"),
        (None, "           2440 and 2441 MHz tell you nothing that 2440 alone did not."),
        (None, "POWER_CAL  Set the generator, read the PEAK amplitude in dBFS, type it in column D."),
        (None, "           At least 6 rows at one gain setting. 15+ gives a usable table."),
        (None, "           Optionally repeat 3-4 frequencies at a second and third gain: the solver"),
        (None, "           then measures how far each gain step deviates from its nominal value."),
        (None, ""),
        (FONT_TITLE, "STAYING IN RANGE"),
        (None, f"Rows reading above {CLIP_DBFS:.0f} dBFS are CLIPPING and are DISCARDED by the solver."),
        (None, f"Rows below {WEAK_DBFS:.0f} dBFS are in the noise and are DISCARDED."),
        (None, "If a row flags either way, change ONLY the generator level in column C and re-read."),
        (None, "The maths stays correct because the offset is computed from what you actually set."),
        (None, ""),
        (FONT_TITLE, "SAFETY"),
        (None, "HackRF One maximum input is -5 dBm. Exceeding it destroys the receiver permanently."),
        (None, "Start low and increase. Confirm the cable is on the RF input, not a clock or TX port."),
        (None, ""),
        (None, "Blank rows are ignored. Delete rows you cannot measure; do not enter guesses."),
    ]
    for i, (font, text) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=2, value=text)
        if font:
            cell.font = font
    ws.sheet_view.showGridLines = False


def _setup_sheet(wb: Workbook, device_type: str, serial: str) -> None:
    ws = wb.create_sheet("SETUP")
    _widths(ws, {1: 34, 2: 22, 3: 74})
    ws["A1"] = "SETUP - fill the yellow cells"
    ws["A1"].font = FONT_TITLE

    rows = [
        ("device_type", device_type or "HACKRF", "HACKRF, USRP, or the profile name used in the app."),
        ("serial", serial or "", "Exact serial string. Blank writes to the device default instead."),
        ("lna_db", 24, "LNA / IF gain. Keep constant for every row in this workbook."),
        ("amp_db", 0, "Front-end amp. 0 = off. Keep constant."),
        ("sample_rate_msps", 20, "Sample rate used while measuring. Affects RBW and noise floor."),
        ("attenuator_db", 10, "External pad AT THE SDR INPUT during calibration. 0 = no pad."),
        ("cable_loss_db", 0, "Measured cable / splitter / coupler loss between generator and pad."),
        ("operational_attenuation_db", 0, "Pad left fitted during NORMAL use. Usually 0. Added back at runtime."),
        ("operator", "", "Who performed this calibration."),
        ("date", datetime.now().strftime("%Y-%m-%d"), "Calibration date."),
        ("generator_model", "", "Signal generator make/model and its own uncertainty if known."),
        ("notes", "", "Anything that would change the result if repeated."),
    ]
    _write_header(ws, 3, ["Field", "Value", "Notes"])
    for i, (key, value, note) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=key).fill = FILL_FIXED
        ws.cell(row=i, column=1).border = BORDER
        cell = ws.cell(row=i, column=2, value=value)
        cell.fill = FILL_INPUT
        cell.border = BORDER
        ws.cell(row=i, column=3, value=note).font = Font(size=9, color="666666")
    ws.sheet_view.showGridLines = False


def _freq_sheet(wb: Workbook, points: list[float]) -> None:
    ws = wb.create_sheet("FREQ_CAL")
    _widths(ws, {1: 16, 2: 18, 3: 14, 4: 14, 5: 40})
    ws["A1"] = "FREQ_CAL - tune to column A, read the marker, type it in column B"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = ("Displayed_MHz is what the application shows, NOT what the generator is set to. "
                "Use the narrowest practical span so one FFT bin is small.")
    ws["A2"].font = Font(size=9, color="666666")

    _write_header(ws, 4, ["Known_MHz", "Displayed_MHz", "Error_Hz", "Error_ppm", "Status"])
    start = 5
    for i, freq in enumerate(points):
        r = start + i
        ws.cell(row=r, column=1, value=freq).fill = FILL_FIXED
        ws.cell(row=r, column=2).fill = FILL_INPUT
        ws.cell(row=r, column=3, value=f'=IF(B{r}="","",(B{r}-A{r})*1000000)').fill = FILL_CALC
        ws.cell(row=r, column=4, value=f'=IF(B{r}="","",ROUND((B{r}-A{r})/A{r}*1000000,3))').fill = FILL_CALC
        ws.cell(row=r, column=5, value=(
            f'=IF(B{r}="","",IF(ABS(C{r})>5000000,"suspicious - check units (MHz not Hz)","ok"))'
        )).fill = FILL_CALC
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BORDER

    last = start + len(points) - 1
    ws.cell(row=last + 2, column=1, value="points entered").font = Font(bold=True)
    ws.cell(row=last + 2, column=2, value=f"=COUNT(B{start}:B{last})")
    ws.cell(row=last + 3, column=1, value="span (MHz)").font = Font(bold=True)
    ws.cell(row=last + 3, column=2, value=(
        f'=IF(COUNT(B{start}:B{last})<2,"enter data",'
        f'MAX(IF(B{start}:B{last}<>"",A{start}:A{last}))-MIN(IF(B{start}:B{last}<>"",A{start}:A{last})))'
    ))
    ws.cell(row=last + 5, column=1, value=(
        "A constant Error_Hz across frequency is a fixed offset. An Error_Hz that grows in "
        "proportion to frequency (constant Error_ppm) is an oscillator error. The solver fits both."
    )).font = Font(size=9, color="666666")
    ws.sheet_view.showGridLines = False


def _power_sheet(wb: Workbook, points: list[float], gain_db: float, siggen_dbm: float) -> None:
    ws = wb.create_sheet("POWER_CAL")
    _widths(ws, {1: 12, 2: 10, 3: 13, 4: 15, 5: 13, 6: 12, 7: 12, 8: 34})
    ws["A1"] = "POWER_CAL - set the generator, read the peak amplitude, type it in column D"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = ("Measured_dBFS is the PEAK MARKER amplitude on the tone, not channel power. "
                "Change column C freely if a row clips or reads too weak.")
    ws["A2"].font = Font(size=9, color="666666")

    _write_header(ws, 4, [
        "Freq_MHz", "Gain_dB", "SigGen_dBm", "Measured_dBFS",
        "InputPwr_dBm", "Offset_dB", "Base_dB", "Status",
    ])
    start = 5
    for i, freq in enumerate(points):
        r = start + i
        ws.cell(row=r, column=1, value=freq).fill = FILL_FIXED
        ws.cell(row=r, column=2, value=gain_db).fill = FILL_INPUT
        ws.cell(row=r, column=3, value=siggen_dbm).fill = FILL_INPUT
        ws.cell(row=r, column=4).fill = FILL_INPUT
        ws.cell(row=r, column=5, value=f"=C{r}-SETUP!$B$9-SETUP!$B$10").fill = FILL_CALC
        ws.cell(row=r, column=6, value=f'=IF(D{r}="","",E{r}-D{r})').fill = FILL_CALC
        ws.cell(row=r, column=7, value=f'=IF(D{r}="","",F{r}+B{r})').fill = FILL_CALC
        ws.cell(row=r, column=8, value=(
            f'=IF(D{r}="","",IF(D{r}>{CLIP_DBFS},"CLIPPING - lower SigGen 10 dB",'
            f'IF(D{r}<{WEAK_DBFS},"too weak - raise SigGen 10 dB","ok")))'
        )).fill = FILL_CALC
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = BORDER

    last = start + len(points) - 1
    ws.cell(row=last + 2, column=1, value="valid points").font = Font(bold=True)
    ws.cell(row=last + 2, column=2, value=f'=COUNTIF(H{start}:H{last},"ok")')
    ws.cell(row=last + 3, column=1, value="ripple (dB)").font = Font(bold=True)
    ws.cell(row=last + 3, column=2, value=(
        f'=IF(COUNT(G{start}:G{last})<4,"enter data",'
        f'ROUND(MAX(G{start}:G{last})-MIN(G{start}:G{last}),2))'
    ))
    ws.cell(row=last + 5, column=1, value=(
        "Ripple here is real front-end response plus any residual mismatch. Several dB of smooth "
        "variation across a wide span is normal and is exactly what the table corrects. Fast "
        "oscillation with a fixed period is a standing wave: fit a bigger pad and repeat."
    )).font = Font(size=9, color="666666")
    ws.cell(row=last + 7, column=1, value=(
        "To calibrate a second gain setting, copy any block of rows, change Gain_dB, and adjust "
        "SigGen_dBm so nothing clips. The solver measures each gain against the busiest one."
    )).font = Font(size=9, color="666666")
    ws.sheet_view.showGridLines = False


def build_template(
    device_type: str = "HACKRF",
    serial: str = "",
    freq_points_mhz: list[float] | None = None,
    power_points_mhz: list[float] | None = None,
    gain_db: float = 20.0,
    siggen_dbm: float = -40.0,
) -> bytes:
    """Return a blank calibration workbook as .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)
    _readme(wb)
    _setup_sheet(wb, device_type, serial)
    _freq_sheet(wb, list(freq_points_mhz or DEFAULT_FREQ_POINTS_MHZ))
    _power_sheet(wb, list(power_points_mhz or DEFAULT_POWER_POINTS_MHZ), gain_db, siggen_dbm)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_filename(device_type: str = "HACKRF", serial: str = "") -> str:
    tag = (serial or "device").replace(" ", "_")[:24]
    return f"calibration_{device_type.upper()}_{tag}_{datetime.now():%Y%m%d}.xlsx"


# ===========================================================================
# Parsing and solving
# ===========================================================================
@dataclass
class CalibrationResult:
    ok: bool = False
    device_type: str = ""
    serial: str = ""
    values: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    summary: list[str] = field(default_factory=list)
    freq_points: int = 0
    power_points: int = 0
    rejected_rows: list[str] = field(default_factory=list)

    def report(self) -> str:
        out = []
        if self.errors:
            out.append("ERRORS")
            out += [f"  - {e}" for e in self.errors]
        if self.summary:
            if out:
                out.append("")
            out.append("RESULT")
            out += [f"  {s}" for s in self.summary]
        if self.warnings:
            out.append("")
            out.append("WARNINGS")
            out += [f"  - {w}" for w in self.warnings]
        if self.rejected_rows:
            out.append("")
            out.append("REJECTED ROWS")
            out += [f"  - {r}" for r in self.rejected_rows]
        return "\n".join(out) or "Nothing to report."


def _num(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value or value.startswith("="):
            return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(result) or math.isinf(result) else result


def _median(values: list[float]) -> float:
    ordered = sorted(values)
    n = len(ordered)
    if n == 0:
        return 0.0
    mid = n // 2
    return ordered[mid] if n % 2 else 0.5 * (ordered[mid - 1] + ordered[mid])


def _lstsq_line(xs: list[float], ys: list[float]) -> tuple[float, float]:
    """Ordinary least squares y = intercept + slope*x, no NumPy dependency."""
    n = len(xs)
    mx = sum(xs) / n
    my = sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    if sxx == 0.0:
        return my, 0.0
    slope = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sxx
    return my - slope * mx, slope


def _read_setup(ws) -> dict:
    setup = {}
    for row in ws.iter_rows(min_row=4, max_col=2):
        key = row[0].value
        if isinstance(key, str) and key.strip():
            setup[key.strip().lower()] = row[1].value
    return setup


def solve_workbook(source) -> CalibrationResult:
    """Parse a filled workbook and derive calibration values.

    ``source`` may be a path, a file-like object, or raw bytes.
    """
    result = CalibrationResult()

    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    try:
        wb = load_workbook(source, data_only=False)
    except Exception as exc:
        result.errors.append(f"Could not open the workbook: {exc}")
        return result

    missing = [s for s in ("SETUP", "FREQ_CAL", "POWER_CAL") if s not in wb.sheetnames]
    if missing:
        result.errors.append(
            f"Missing sheet(s): {', '.join(missing)}. Download a fresh template and refill it."
        )
        return result

    setup = _read_setup(wb["SETUP"])
    result.device_type = str(setup.get("device_type") or "HACKRF").strip().upper()
    result.serial = str(setup.get("serial") or "").strip()

    atten = _num(setup.get("attenuator_db")) or 0.0
    cable = _num(setup.get("cable_loss_db")) or 0.0
    operational = _num(setup.get("operational_attenuation_db")) or 0.0

    if atten == 0.0:
        result.warnings.append(
            "attenuator_db is 0. Without a pad at the SDR input, cable standing waves are baked "
            "into the table and it will be wrong by up to +/-2 dB with any other source."
        )

    values: dict = {}

    # --- Frequency ---------------------------------------------------------
    freqs_mhz, errors_hz = [], []
    for row in wb["FREQ_CAL"].iter_rows(min_row=5, max_col=2):
        known = _num(row[0].value)
        displayed = _num(row[1].value)
        if known is None or displayed is None or known <= 0:
            continue
        error_hz = (displayed - known) * 1e6
        if abs(error_hz) > 5e6:
            result.rejected_rows.append(
                f"FREQ_CAL {known:g} MHz: error {error_hz/1e6:.2f} MHz is implausible "
                "(units mixed up?) - discarded"
            )
            continue
        freqs_mhz.append(known)
        errors_hz.append(error_hz)

    result.freq_points = len(freqs_mhz)
    if result.freq_points == 0:
        result.warnings.append("No FREQ_CAL rows filled. Frequency calibration left unchanged.")
    elif result.freq_points == 1:
        fixed = errors_hz[0]
        values["frequency_axis_offset_hz"] = round(-fixed, 1)
        values["ppm_offset"] = 0.0
        values["frequency_offset_table"] = []
        result.warnings.append(
            "Only one FREQ_CAL point. A fixed offset was applied, but one point cannot "
            "distinguish a constant offset from an oscillator ppm error. Add points at widely "
            "separated frequencies."
        )
        result.summary.append(f"Frequency: fixed offset {-fixed:+.0f} Hz (single point, unverified)")
    else:
        span_mhz = max(freqs_mhz) - min(freqs_mhz)
        fixed, ppm = _lstsq_line(freqs_mhz, errors_hz)
        residuals = [e - (fixed + ppm * f) for f, e in zip(freqs_mhz, errors_hz)]
        max_resid = max(abs(r) for r in residuals)
        ptp = max(errors_hz) - min(errors_hz)

        if span_mhz < 0.05 * max(freqs_mhz):
            fixed = _median(errors_hz)
            ppm = 0.0
            residuals = [e - fixed for e in errors_hz]
            max_resid = max(abs(r) for r in residuals)
            result.warnings.append(
                f"FREQ_CAL points span only {span_mhz:.0f} MHz. That is too narrow to separate a "
                "fixed offset from ppm, so a fixed offset alone was fitted. Spread the points."
            )

        values["frequency_axis_offset_hz"] = round(-fixed, 1)
        values["ppm_offset"] = round(ppm, 4)
        result.summary.append(
            f"Frequency: fixed {-fixed:+.0f} Hz, oscillator {ppm:+.3f} ppm "
            f"(from {result.freq_points} points over {span_mhz:.0f} MHz)"
        )

        threshold = max(FREQ_RESIDUAL_FLOOR_HZ, FREQ_RESIDUAL_FRACTION * ptp)
        if max_resid > threshold and result.freq_points >= 4:
            values["frequency_offset_table"] = [
                {"freq_hz": round(f * 1e6), "offset_hz": round(-r, 1)}
                for f, r in sorted(zip(freqs_mhz, residuals))
            ]
            result.warnings.append(
                f"Frequency error is not a clean straight line: worst residual {max_resid:.0f} Hz "
                f"against a {threshold:.0f} Hz threshold. A residual interpolation table was "
                "written alongside the linear fit to absorb it."
            )
        else:
            values["frequency_offset_table"] = []
            result.summary.append(
                f"Frequency fit residual: {max_resid:.0f} Hz worst case - linear model is adequate"
            )

    # --- Power -------------------------------------------------------------
    samples: list[tuple[float, float, float]] = []   # (freq_hz, gain_db, base_db)
    for row in wb["POWER_CAL"].iter_rows(min_row=5, max_col=4):
        freq = _num(row[0].value)
        gain = _num(row[1].value)
        siggen = _num(row[2].value)
        measured = _num(row[3].value)
        if freq is None or measured is None or siggen is None:
            continue
        if gain is None:
            result.rejected_rows.append(f"POWER_CAL {freq:g} MHz: Gain_dB is blank - discarded")
            continue
        if measured > CLIP_DBFS:
            result.rejected_rows.append(
                f"POWER_CAL {freq:g} MHz gain {gain:g}: {measured:.2f} dBFS is clipping - discarded"
            )
            continue
        if measured < WEAK_DBFS:
            result.rejected_rows.append(
                f"POWER_CAL {freq:g} MHz gain {gain:g}: {measured:.2f} dBFS is in the noise - discarded"
            )
            continue
        base = (siggen - atten - cable - measured) + gain
        samples.append((freq * 1e6, gain, base))

    result.power_points = len(samples)
    if result.power_points < 3:
        result.warnings.append(
            f"Only {result.power_points} usable POWER_CAL rows. Power calibration left unchanged; "
            "at least 3 (realistically 10+) are needed for a usable table."
        )
    else:
        by_gain: dict[float, list[tuple[float, float]]] = {}
        for freq_hz, gain, base in samples:
            by_gain.setdefault(gain, []).append((freq_hz, base))

        ref_gain = max(by_gain, key=lambda g: (len({f for f, _ in by_gain[g]}), -abs(g)))
        ref_rows = by_gain[ref_gain]

        grouped: dict[float, list[float]] = {}
        for freq_hz, base in ref_rows:
            grouped.setdefault(freq_hz, []).append(base)

        base_table = []
        for freq_hz in sorted(grouped):
            bases = grouped[freq_hz]
            if len(bases) > 1:
                scatter = max(bases) - min(bases)
                if scatter > POWER_SCATTER_WARN_DB:
                    result.warnings.append(
                        f"Repeat readings at {freq_hz/1e6:.0f} MHz disagree by {scatter:.1f} dB. "
                        "The median was used; check for drift or a loose connector."
                    )
            base_table.append({"freq_hz": round(freq_hz), "base_db": round(_median(bases), 2)})

        values["power_base_table"] = base_table
        values["power_reference_gain_db"] = ref_gain
        values["external_attenuation_db"] = operational
        values["power_offset_db"] = None   # table supersedes the legacy scalar

        ripple = max(p["base_db"] for p in base_table) - min(p["base_db"] for p in base_table)
        result.summary.append(
            f"Power: {len(base_table)} frequency points at gain {ref_gain:g} dB, "
            f"{base_table[0]['freq_hz']/1e6:.0f}-{base_table[-1]['freq_hz']/1e6:.0f} MHz, "
            f"{ripple:.1f} dB of response variation"
        )

        ref_lookup = {p["freq_hz"]: p["base_db"] for p in base_table}
        gain_table = [{"gain_db": ref_gain, "delta_db": 0.0}]
        for gain in sorted(by_gain):
            if gain == ref_gain:
                continue
            deltas = []
            for freq_hz, base in by_gain[gain]:
                key = round(freq_hz)
                if key in ref_lookup:
                    deltas.append(base - ref_lookup[key])
            if not deltas:
                result.warnings.append(
                    f"Gain {gain:g} dB shares no frequency with gain {ref_gain:g} dB, so its "
                    "deviation could not be measured. Repeat at least one common frequency."
                )
                continue
            delta = round(_median(deltas), 2)
            gain_table.append({"gain_db": gain, "delta_db": delta})
            verdict = "nominal" if abs(delta) < 1.0 else f"DEVIATES by {delta:+.2f} dB"
            result.summary.append(f"Gain {gain:g} dB: {verdict}")

        values["power_gain_table"] = sorted(gain_table, key=lambda p: p["gain_db"])
        if len(gain_table) == 1:
            result.warnings.append(
                f"Only gain {ref_gain:g} dB was calibrated. Readings at other gains assume the "
                "gain steps are exactly nominal, which is often wrong by several dB at the "
                "extremes. Add rows at the gains you actually use."
            )

    if not values:
        result.errors.append("Nothing usable was found in the workbook. No changes were written.")
        return result

    values["metadata"] = {
        "calibrated_at": datetime.now().isoformat(timespec="seconds"),
        "operator": str(setup.get("operator") or ""),
        "generator_model": str(setup.get("generator_model") or ""),
        "lna_db": _num(setup.get("lna_db")),
        "amp_db": _num(setup.get("amp_db")),
        "sample_rate_msps": _num(setup.get("sample_rate_msps")),
        "attenuator_db": atten,
        "cable_loss_db": cable,
        "notes": str(setup.get("notes") or ""),
        "freq_points": result.freq_points,
        "power_points": result.power_points,
        "rejected_rows": len(result.rejected_rows),
    }

    result.values = values
    result.ok = True
    return result


def apply_workbook(source, path: Path | None = None) -> CalibrationResult:
    """Solve a workbook and, if valid, write the result into calibration.json."""
    from .calibration import save_device_calibration

    result = solve_workbook(source)
    if not result.ok:
        return result
    try:
        target = save_device_calibration(
            result.device_type, result.serial, result.values, path
        )
    except OSError as exc:
        result.ok = False
        result.errors.append(f"Could not write the calibration file: {exc}")
        return result

    scope = f"serial {result.serial}" if result.serial else "device default"
    result.summary.append(f"Written to {target} under {result.device_type} / {scope}")
    result.summary.append("Restart acquisition for the new calibration to take effect.")
    return result